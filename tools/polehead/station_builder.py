#!/usr/bin/env python3
"""Read sheet data (crew notes, notes, GPS, elevation) from a Station Builder xlsx.

    python station_builder.py "<TD> - STATION BUILDER.xlsx" 1414703E 4904169E ... > sheets.json

Reads ONLY columns N:Q, rows 4 to 46 of each pole tab (the tabs are thousands of
rows long). Output: {pole: {crew, notes, gps, elev}} plus a "report" list of
things worth a look. Merge these into the job spec's poles.
"""
from __future__ import annotations

import json
import re
import sys

FIX = {'DEFUALT': 'DEFAULT', 'REQUIRMENT': 'REQUIREMENT', 'EASMENT': 'EASEMENT', 'UKNOWN': 'UNKNOWN', 'VEGITATION': 'VEGETATION',
       'OBSTUCTION': 'OBSTRUCTION', 'ISNTALLED': 'INSTALLED'}


def clean(t):
    t = ' '.join(str(t).replace('\n', ' ').split()).upper()
    for a, b in FIX.items():
        t = t.replace(a, b)
    t = re.sub(r'\bC0MM\b', 'COMM', t)          # zero typed for the letter O
    return t


def read_rows(xlsx, pole):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    if pole in wb.sheetnames:
        ws = wb[pole]
    else:
        # tabs are sometimes suffixed, e.g. "1296855E - EXTREME" or "1754156E EXTREME"
        hits = [n for n in wb.sheetnames if n.split()[0].upper() == pole.upper()]
        if len(hits) != 1:
            raise KeyError(pole)
        ws = wb[hits[0]]
    return [tuple(c.value for c in row) for row in ws.iter_rows(min_row=4, max_row=46, min_col=14, max_col=17)]


def sheet_from_rows(pole, rows, report):
    """rows[i] = (N, O, P, Q) of row 4 + i. Fill-ins can sit in O, P or Q (planners vary)."""
    crew, notes = [], []
    for i in range(13, 22):                      # N17..N25
        v = rows[i][0]
        if v:
            c = clean(v)
            crew.append(c if c.startswith('-') else '-' + c)
    for i in range(23, 43):                      # N27..N46
        v, o, p, q = (tuple(rows[i]) + (None,) * 4)[:4]
        o, p, q = [x if not isinstance(x, str) or x.strip() else None for x in (o, p, q)]
        if not v and o:                          # planners type whole notes in column O with N empty: still a note
            v, o, p, q = o, None, None, None
        if not v:
            continue
        t = clean(v)
        fill = ' '.join(str(x).strip() for x in (o, p, q) if x)
        if 'POLE SET IN' in t:
            if fill:
                f = clean(fill)
                if 'INSP.' in f:
                    a, _, b = f.partition('INSP.')
                else:                            # "2014 , N/A" = set year, inspection year
                    f = re.sub(r"^(\d{4})'\s+", r'\1, ', f)   # "2019' N/A": apostrophe typed for the comma
                    a, _, b = f.partition(',')
                na = lambda x: 'N/A' if x.strip(' ,') in ('', 'NA', 'N/A') else x.strip(' ,')
                t = '-POLE SET IN %s, INSP. %s' % (na(a), na(b))
            else:
                t = '-POLE SET IN N/A, INSP. N/A'
        elif 'STREET NAME' in t:
            if fill and 'ACCESS' in clean(fill):   # whole access note typed as the fill-in: it replaces the template
                t = clean(fill)
            elif fill:
                t = t.replace('STREET NAME', clean(fill))
        elif '___' in t and fill:
            blanks = re.findall(r'_+(?:/_+)*', t)
            parts = [x.strip() for x in clean(fill).split(',')]
            if len(blanks) > 1 and len(parts) == len(blanks):
                # several blanks, comma separated fill-ins: fill them in order
                # ("NOTIFICATION #____ REQUESTED ON __/__/__" + "415588905, 9/3/2026")
                for x in parts:
                    t = re.sub(r'_+(?:/_+)*', lambda m, x=x: x, t, count=1)
            else:
                t = t.split('___')[0].rstrip() + ' ' + clean(fill)
        if not t.startswith('-'):
            t = '-' + t
        t = re.sub(r'\bVIA VIA\b', 'VIA', t)              # fill-in that repeats the template word
        t = re.sub(r'\b(OF|AT|R/O)(\d)', r'\1 \2', t)      # missing space before a house number (OF1211)
        if t in notes:
            continue                             # duplicates are dropped silently
        notes.append(t)
    gps = [str(rows[5][1] or ''), str(rows[7][1] or '')]
    elev = str(rows[9][1] or '').upper().replace("'", '').replace('FT', '').strip()
    if not gps[0] or not gps[1]:
        report.append('%s: GPS missing in Station Builder' % pole)
    return dict(crew=crew, notes=notes, gps=gps, elev=elev)


def main(argv):
    if len(argv) < 2:
        print(__doc__, file=sys.stderr)
        return 2
    xlsx, poles = argv[0], argv[1:]
    out, report = {}, []
    for p in poles:
        try:
            out[p] = sheet_from_rows(p, read_rows(xlsx, p), report)
        except KeyError:
            report.append('%s: no Station Builder tab' % p)
    out['report'] = report
    json.dump(out, sys.stdout, indent=1)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
